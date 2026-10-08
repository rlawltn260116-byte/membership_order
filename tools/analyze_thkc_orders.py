from __future__ import annotations

import json
import re
import unicodedata
import argparse
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT_ROOT / "worker" / "order-runner" / "data" / "thkc-learned-catalog.json"
LOCAL_OUTPUT = PROJECT_ROOT / ".local-analysis" / "thkc-analysis.json"
START_DATE = 20260803
END_DATE = 20261007

HEADERS = {
    "주문번호": "orderId",
    "납품처": "destination",
    "출고처": "destination",
    "납품처전화번호": "phone",
    "출고처전화번호": "phone",
    "품명": "product",
    "품목": "product",
    "주문수량": "quantity",
    "비고(내역)": "note",
    "비고(업체)": "note",
}

CATEGORY_PREFIXES = (
    "미끄럼방지매트리스",
    "미끄럼방지양말",
    "성인용보행기",
    "자세변환용구",
    "욕창예방방석",
    "이동변기부품",
    "휠체어부품",
    "수동휠체어",
    "안전손잡이",
    "지팡이",
)


def text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip()


def parse_file_date(name: str) -> int | None:
    match = re.match(r"^(\d{8})_", name)
    return int(match.group(1)) if match else None


def shift_of(name: str) -> str:
    if "오전" in name:
        return "오전"
    if "오후" in name:
        return "오후"
    return "미분류"


def split_destination(value: str) -> tuple[str, str]:
    if "_" not in value:
        spaced_slash = re.search(r"\s+/\s*([^/]+)$", value)
        if spaced_slash:
            return value[: spaced_slash.start()].strip(), spaced_slash.group(1).strip()
        return value, ""
    address, recipient = value.rsplit("_", 1)
    return address.strip(), recipient.strip()


def split_product(value: str) -> tuple[str, str]:
    value = text(value)
    depth = 0
    for index, char in enumerate(value):
        if char == "(":
            depth += 1
        elif char == ")" and depth:
            depth -= 1
        elif char == "/" and depth == 0:
            return value[:index].strip(), value[index + 1 :].strip()
    compact_option = re.fullmatch(r"(DA-006)\((초록|회색)\)", value)
    if compact_option:
        return compact_option.group(1), compact_option.group(2)
    return value, ""


def model_of(base: str) -> str:
    if base == "DA-006":
        return base
    for prefix in CATEGORY_PREFIXES:
        if base.startswith(prefix):
            return base[len(prefix):].strip()
    return base


def category_of(base: str, model: str) -> str:
    if model == "DA-006":
        return "미끄럼방지매트리스"
    for prefix in CATEGORY_PREFIXES:
        if base.startswith(prefix):
            return prefix
    return ""


def find_header_blocks(sheet) -> list[tuple[int, dict[str, int]]]:
    blocks: list[tuple[int, dict[str, int]]] = []
    for row_no in range(1, min(sheet.max_row, 25) + 1):
        mapping: dict[str, int] = {}
        for col in range(1, sheet.max_column + 1):
            canonical = HEADERS.get(text(sheet.cell(row_no, col).value))
            if canonical:
                mapping[canonical] = col
        if len(mapping) >= 4 and {"destination", "phone", "product", "quantity"}.issubset(mapping):
            blocks.append((row_no, mapping))
    return blocks


def order_type(order_id: str) -> str:
    if re.fullmatch(r"SO\d{10}", order_id):
        return "CUSTOMER_ORDER"
    if not order_id or order_id.upper() == "THK":
        return "STOCK_OR_BRANCH_ORDER"
    return "UNKNOWN"


def main(source_dir: Path) -> None:
    files = []
    for path in sorted(source_dir.glob("*.xlsx")):
        day = parse_file_date(path.name)
        if day is not None and START_DATE <= day <= END_DATE:
            files.append(path)

    records: list[dict[str, object]] = []
    file_summaries: list[dict[str, object]] = []
    header_patterns: Counter[tuple[str, ...]] = Counter()
    sheet_names: Counter[str] = Counter()
    workbook_errors: list[dict[str, str]] = []

    for path in files:
        try:
            workbook = load_workbook(path, data_only=True, read_only=True)
            sheet = workbook[workbook.sheetnames[0]]
            sheet_names[text(sheet.title)] += 1
            header_blocks = find_header_blocks(sheet)
            if not header_blocks:
                raise ValueError("지원되는 주문 헤더를 찾지 못했습니다.")
            for _, mapping in header_blocks:
                header_pattern = tuple(sorted(mapping, key=lambda key: mapping[key]))
                header_patterns[header_pattern] += 1
            used_rows = 0
            active_mapping: dict[str, int] | None = None
            blocks_by_row = dict(header_blocks)
            for row_no in range(1, sheet.max_row + 1):
                if row_no in blocks_by_row:
                    active_mapping = blocks_by_row[row_no]
                    continue
                if active_mapping is None:
                    continue
                values = {
                    key: sheet.cell(row_no, col).value
                    for key, col in active_mapping.items()
                }
                if not any(values.get(key) not in (None, "") for key in ("orderId", "destination", "phone", "product", "quantity")):
                    continue
                used_rows += 1
                order_id = text(values.get("orderId"))
                destination = text(values.get("destination"))
                phone = text(values.get("phone"))
                product = text(values.get("product"))
                quantity = text(values.get("quantity"))
                note = text(values.get("note"))
                address, recipient = split_destination(destination)
                base, option = split_product(product)
                model = model_of(base)
                category = category_of(base, model)
                try:
                    quantity_number: int | float | str = float(quantity)
                    if quantity_number.is_integer():
                        quantity_number = int(quantity_number)
                except (TypeError, ValueError):
                    quantity_number = quantity
                records.append(
                    {
                        "file": path.name,
                        "date": parse_file_date(path.name),
                        "shift": shift_of(path.name),
                        "row": row_no,
                        "orderType": order_type(order_id),
                        "orderId": order_id,
                        "destinationRaw": destination,
                        "address": address,
                        "recipient": recipient,
                        "phone": phone,
                        "productRaw": product,
                        "productBase": base,
                        "category": category,
                        "model": model,
                        "option": option,
                        "quantity": quantity_number,
                        "note": note,
                    }
                )
            file_summaries.append(
                {
                    "file": path.name,
                    "date": parse_file_date(path.name),
                    "shift": shift_of(path.name),
                    "sheet": text(sheet.title),
                    "rows": used_rows,
                    "headerBlocks": [
                        {"row": row, "columns": mapping}
                        for row, mapping in header_blocks
                    ],
                    "specialFilename": any(token in path.name for token in ("추가포함", "미접수주문", "최종", "발주서2")),
                }
            )
            workbook.close()
        except Exception as exc:  # preserve analysis coverage rather than stopping at one file
            workbook_errors.append({"file": path.name, "error": str(exc)})

    by_product: dict[str, list[dict[str, object]]] = defaultdict(list)
    for record in records:
        by_product[str(record["productRaw"])].append(record)

    product_catalog = []
    for raw, rows in sorted(by_product.items(), key=lambda item: (-len(item[1]), item[0])):
        quantities = [r["quantity"] for r in rows if isinstance(r["quantity"], (int, float))]
        type_counts = Counter(str(r["orderType"]) for r in rows)
        product_catalog.append(
            {
                "rawName": raw,
                "baseName": rows[0]["productBase"],
                "category": rows[0]["category"],
                "model": rows[0]["model"],
                "option": rows[0]["option"],
                "orderRows": len(rows),
                "totalQuantity": sum(quantities),
                "customerOrderRows": type_counts["CUSTOMER_ORDER"],
                "stockOrBranchOrderRows": type_counts["STOCK_OR_BRANCH_ORDER"],
                "firstSeen": min(int(r["date"]) for r in rows),
                "lastSeen": max(int(r["date"]) for r in rows),
            }
        )

    order_ids = [str(r["orderId"]) for r in records if r["orderId"]]
    customer_records = [r for r in records if r["orderType"] == "CUSTOMER_ORDER"]
    customer_id_counts = Counter(str(r["orderId"]) for r in customer_records)
    stock_records = [r for r in records if r["orderType"] == "STOCK_OR_BRANCH_ORDER"]
    unknown_records = [r for r in records if r["orderType"] == "UNKNOWN"]
    duplicate_ids = {
        key: count for key, count in Counter(order_ids).items() if count > 1
    }
    notes = Counter(str(r["note"]) for r in records if r["note"])
    invalid_phone = [
        r
        for r in records
        if r["phone"] and not re.fullmatch(r"0\d{1,2}-\d{3,4}-\d{4}", str(r["phone"]))
    ]
    missing_address = [r for r in records if not r["address"]]
    missing_recipient = [r for r in records if not r["recipient"]]
    missing_order_id = [r for r in records if not r["orderId"]]
    missing_phone = [r for r in records if not r["phone"]]
    nonpositive_qty = [r for r in records if not isinstance(r["quantity"], (int, float)) or r["quantity"] <= 0]

    base_groups: list[dict[str, object]] = []
    grouped: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for product in product_catalog:
        grouped[(str(product["category"]), str(product["model"]))].append(product)
    for (category, model), variants in sorted(grouped.items(), key=lambda item: (-sum(int(v["orderRows"]) for v in item[1]), item[0])):
        base_groups.append(
            {
                "category": category,
                "model": model,
                "aliases": sorted({str(v["baseName"]) for v in variants}),
                "options": sorted({str(v["option"]) for v in variants if v["option"]}),
                "orderRows": sum(int(v["orderRows"]) for v in variants),
                "totalQuantity": sum(float(v["totalQuantity"]) for v in variants),
            }
        )

    payload = {
        "generatedAt": datetime.now().astimezone().isoformat(),
        "source": {
            "company": "티에이치케이컴퍼니",
            "code": "THKC",
            "channel": "이로움",
            "directory": str(source_dir),
            "dateRange": {"from": START_DATE, "to": END_DATE},
        },
        "summary": {
            "files": len(files),
            "filesParsed": len(file_summaries),
            "rows": len(records),
            "distinctOrderIds": len(set(order_ids)),
            "distinctProductStrings": len(by_product),
            "distinctProductBases": len({str(r["productBase"]) for r in records}),
            "morningFiles": sum(1 for f in file_summaries if f["shift"] == "오전"),
            "afternoonFiles": sum(1 for f in file_summaries if f["shift"] == "오후"),
            "specialFiles": sum(1 for f in file_summaries if f["specialFilename"]),
            "workbookErrors": len(workbook_errors),
            "customerOrderRows": len(customer_records),
            "customerOrders": len(customer_id_counts),
            "stockOrBranchRows": len(stock_records),
            "unknownOrderRows": len(unknown_records),
            "multiLineCustomerOrders": sum(1 for count in customer_id_counts.values() if count > 1),
        },
        "schema": {
            "sheetNames": dict(sheet_names),
            "headerPatterns": [
                {"canonicalColumns": list(headers), "blocks": count}
                for headers, count in header_patterns.most_common()
            ],
            "columns": ["주문번호", "납품처", "납품처전화번호", "품명", "주문수량", "비고(내역)"],
            "destinationRule": "납품처는 주로 주소_수취인 형식이며 마지막 밑줄로 분리. 밑줄이 없을 때 공백을 동반한 마지막 슬래시를 보조 구분자로 사용",
            "postcodePresent": False,
        },
        "quality": {
            "missingOrderIdRows": len(missing_order_id),
            "missingRecipientRows": len(missing_recipient),
            "missingAddressRows": len(missing_address),
            "missingPhoneRows": len(missing_phone),
            "invalidPhoneRows": len(invalid_phone),
            "nonpositiveOrNonNumericQuantityRows": len(nonpositive_qty),
            "duplicateOrderIds": duplicate_ids,
            "notes": [{"text": note, "rows": count} for note, count in notes.most_common()],
            "examples": {
                "missingOrderId": missing_order_id[:20],
                "missingRecipient": missing_recipient[:20],
                "invalidPhone": invalid_phone[:20],
                "nonpositiveQuantity": nonpositive_qty[:20],
            },
        },
        "products": product_catalog,
        "productGroups": base_groups,
        "files": file_summaries,
        "records": records,
        "errors": workbook_errors,
    }

    sanitized = {
        "generatedAt": payload["generatedAt"],
        "source": {
            "company": "티에이치케이컴퍼니",
            "code": "THKC",
            "channel": "이로움",
            "dateRange": payload["source"]["dateRange"],
        },
        "summary": payload["summary"],
        "schema": payload["schema"],
        "orderRules": {
            "customerOrderIdPattern": "^SO\\d{10}$",
            "blankOrThkOrderId": "지점/재고 주문 후보. 고객 주문과 분리해 검토",
            "postcode": "원본에 우편번호가 없어 주소 검색 또는 주소 API 보완 필요",
            "recipient": "구분자로 수취인을 얻지 못하면 자동 제출 금지 후 검토",
            "phone": "휴대전화와 지역번호 유선전화를 모두 허용. 숫자만 있는 9~11자리는 표준 하이픈 형식으로 정규화하고 그 외 형식은 검토",
            "quantity": "양의 숫자만 실행 가능",
            "notes": "배송지시, 센터 경유, 출입정보가 섞이므로 주문 메모로 보존하고 외부 노출 금지",
            "finalSubmission": "운영 페이지 검토 수락 전에는 실행 금지",
        },
        "qualityCounts": {
            "missingOrderIdRows": len(missing_order_id),
            "missingRecipientRows": len(missing_recipient),
            "missingAddressRows": len(missing_address),
            "missingPhoneRows": len(missing_phone),
            "phoneFormatExceptionRows": len(invalid_phone),
            "nonpositiveOrNonNumericQuantityRows": len(nonpositive_qty),
            "multiLineOrderIdCount": len(duplicate_ids),
        },
        "products": product_catalog,
        "productGroups": base_groups,
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(sanitized, ensure_ascii=False, indent=2), encoding="utf-8")
    LOCAL_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    LOCAL_OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "output": str(OUTPUT),
                "localOutput": str(LOCAL_OUTPUT),
                "summary": payload["summary"],
                "quality": {key: value for key, value in payload["quality"].items() if key not in {"examples", "notes", "duplicateOrderIds"}},
                "topProducts": product_catalog[:20],
                "notes": payload["quality"]["notes"][:15],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="THKC 이로움 발주서 구조와 상품 표기를 분석합니다.")
    parser.add_argument("source_dir", type=Path, help="THKC 발주서 .xlsx 파일 폴더")
    args = parser.parse_args()
    main(args.source_dir)
